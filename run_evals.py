"""Eval runner for ShopMate (assignment Step 5).

Reads the test cases from eval_set.xlsx (sheet "Eval Set"), runs each one against the real agent on a
fresh copy of store.db, checks the tools, filters and reply, times it, and writes a scored "Run N" sheet
into the same workbook. The scores are proposals: read the replies, change any score you disagree with,
and fill in "Your verdict".

Run from this folder:
    .venv\\Scripts\\python.exe run_evals.py                   all test cases
    .venv\\Scripts\\python.exe run_evals.py --only E03,E10    some of them
    .venv\\Scripts\\python.exe run_evals.py --pause 30        seconds to wait between test cases (default 20)
"""

import argparse
import datetime
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import time

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(HERE)
sys.path.insert(0, HERE)

import guardrails  # noqa: E402
import tools  # noqa: E402
from agent import MODEL, ShopMateAgent  # noqa: E402

WORKBOOK = os.path.join(HERE, "eval_set.xlsx")
STORE_DB = os.path.join(HERE, "store.db")

# How the runner sets up and checks each test case. The readable version of each case is in the Eval Set sheet.
#   text             what to send, when it differs from shopper_says in the sheet (e.g. a photo with no text)
#   image            photo sent with the message
#   setup_say        messages sent first, to put a product list on screen
#   setup_shown      the (ID:X) list the setup must produce
#   setup_prefs      rows put in the preferences table before the test
#   setup_order      product ID ordered before the test, so order history is not empty
#   latency          which latency target applies: "off_topic", "text" or "photo"
#   tools            tools that should be called, in order ([] = none, None = not scored)
#   forbid_tools     tools that must not be called
#   args             expected tool arguments: text = "contains", "!text" = "must not be", number = "equals"
#   result           expected values in a tool's result ("a.b" reaches into nested values)
#   expect_ids       products that must appear in the reply
#   forbid_ids       products that must not appear in the reply
#   list_format      the reply should be a product list in the PRD format
#   exact            text the reply must contain word for word (REDIRECT and NOT_SOLD are filled in from the store)
#   any_of           the reply must contain at least one of these (not case-sensitive)
#   orders           how many new orders the test should create
#   ordered_id       the product that should be ordered
#   history_order    the reply must show this seeded order's product name and price
#   prefs_after      rows the preferences table must have afterwards
#   off_topic        the off-topic check should stop the message before the agent runs
HONEYS = {1, 2, 3, 4, 5, 6, 7, 8}
CHECKS = {
    "E01": dict(latency="text", tools=["search_products", "get_rating"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "honey", "is_organic": "yes", "max_price": 20}},
                expect_ids={1, 5, 7}, forbid_ids={2, 3, 4, 6, 8, 25}, list_format=True, orders=0),
    "E02": dict(latency="text", tools=["search_products", "get_rating"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "olive"}},
                expect_ids={9}, list_format=True, orders=0),
    "E03": dict(latency="text", setup_say=["cheapest oat milk"], setup_shown=[30],
                tools=["checkout"], args={"checkout": {"product_id": 30}},
                orders=1, ordered_id=30),
    "E04": dict(latency="text", setup_order=1, tools=["get_order_history"], forbid_tools=["checkout"],
                history_order=1, orders=0),
    "E05": dict(latency="off_topic", off_topic=True, tools=[], exact=["REDIRECT"], orders=0),
    "E06": dict(latency="text", tools=["search_products", "get_rating"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "honey"}},
                expect_ids=HONEYS, forbid_ids={25}, list_format=True, orders=0,
                any_of=["which", "number", "would you like"]),
    "E07": dict(latency="text", tools=["search_products"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "decaf"}},
                exact=["NOT_SOLD"], forbid_ids=set(range(1, 33)), orders=0),
    "E08": dict(latency="photo", text="find this", image="honey.png",
                tools=["describe_product_image", "search_products", "get_rating"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "honey", "is_organic": "!yes"}},
                expect_ids=HONEYS, forbid_ids={25}, list_format=True, orders=0),
    "E09": dict(latency="off_topic", text="", image="elephant.png", off_topic=True,
                tools=[], exact=["REDIRECT"], orders=0),
    "E10": dict(latency="text", setup_say=["olive oil with a rating above 4.5"], setup_shown=[9],
                tools=[], forbid_tools=["checkout"], orders=0, any_of=["confirm", "yes"]),
    "E11": dict(latency="text", setup_prefs={"is_organic": "1"},
                tools=["search_products", "get_rating"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "honey"}},
                result={"search_products": {"filters_applied.organic_only": True}},
                expect_ids={1, 3, 5, 7}, forbid_ids={2, 4, 6, 8, 25}, list_format=True, orders=0),
    "E12": dict(latency="text", tools=None, forbid_tools=["checkout"],
                prefs_after={"is_organic": "1"}, orders=0),
    "E13": dict(latency="text", setup_say=["organic honey with at least a 4.7 rating"], setup_shown=[3, 7],
                tools=[], forbid_tools=["checkout"], orders=0, any_of=["#1", "#2"]),
    "E14": dict(latency="text", pause_before=60, tools=["search_products", "get_rating"], forbid_tools=["checkout"],
                args={"search_products": {"keyword": "honey", "is_organic": "yes", "max_price": 20}},
                expect_ids={1, 5, 7}, forbid_ids={2, 3, 4, 6, 8, 25}, list_format=True, orders=0),
}

LIST_LINE = re.compile(r"^#(\d+)\. (.+) \(ID:(\d+)\) — \$(\d+\.\d{2}) ★(\d+(?:\.\d+)?)( — organic)?$")
ONE_PRODUCT_QUESTION = "Would you like to order it? Just say yes or give me the number."

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
YOU_FILL = PatternFill("solid", fgColor="FFF2CC")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
GREY_FILL = PatternFill("solid", fgColor="D9D9D9")
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


# --- Rubric sheet ---------------------------------------------------------------------------

PASS_THRESHOLD_CELL = "Rubric!$B$17"
LATENCY_LABELS = {
    "off_topic": "Latency target: off-topic block (seconds)",
    "text": "Latency target: text search or order (seconds)",
    "photo": "Latency target: photo search (seconds)",
}


def ensure_rubric(wb):
    """Create the Rubric sheet if the workbook doesn't have one yet."""
    if "Rubric" in wb.sheetnames:
        return
    ws = wb.create_sheet("Rubric")
    for col, width in zip("ABCDEF", (36, 30, 30, 30, 26, 46)):
        ws.column_dimensions[col].width = width

    def put(row, values, font=FONT, fill=None):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=value)
            cell.font, cell.alignment = font, WRAP
            if fill:
                cell.fill = fill

    ws["A1"] = "ShopMate eval rubric"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A2"] = ("Score every test case on the checks that apply to it and mark the rest N/A. The runner proposes the scores; "
                "change any score you disagree with, then give your verdict. Points, % and the auto result update by themselves.")
    ws["A2"].font, ws["A2"].alignment = FONT, WRAP
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 30

    put(4, ["Check", "2 points", "1 point", "0 points", "N/A when", "How the runner scores it"], HEADER_FONT, HEADER_FILL)
    put(5, ["Tool", "Right tools, in the right order, and no forbidden tool", "Right tools plus an unneeded extra call",
            "Wrong or missing tool, or a forbidden tool was called (e.g. checkout)", "The test is about an outcome, not a tool",
            "Compares the tools called with the expected list"])
    put(6, ["Filters", "Every filter exactly as expected", "One filter wrong or missing", "Several filters wrong, or the tool wasn't called",
            "No filters to check", "Compares the tool arguments (keyword, max_price, is_organic, product_id) and key result values"])
    put(7, ["Reply content", "The right products or message, and nothing wrong", "Right products or outcome, with a minor issue (e.g. it doesn't ask the expected follow-up)",
            "Wrong or missing products, wrong message, wrong order, or a preference not saved", "Never",
            "Checks the product IDs shown, exact messages, orders created and preferences saved"])
    put(8, ["Format", "Exact list format, (ID:X) on every line, one blank line between entries, the single-product question when one product",
            "One small format slip", "Several slips, or (ID:X) missing", "The reply is not a product list",
            "Checks each list line against: #1. Name (ID:1) — $14.99 ★4.62 — organic"])
    put(9, ["Latency", "Within the target below", "Up to twice the target", "More than twice the target", "Never",
            "Times the reply from sending the message to the full answer, minus any wait for Groq's per-minute limit (shown in its own column)"])

    put(11, ["Critical failure (fails the test case whatever the score)", "How the runner detects it"], HEADER_FONT, HEADER_FILL)
    ws.merge_cells("B11:F11")
    put(12, ["An order placed without a clear yes, or for the wrong product", "A new row in orders when none was expected, or for a different product"])
    put(13, ["A product, price or rating that isn't in the store", "Any (ID:X) line whose name, price or rating doesn't match the store data"])
    put(14, ["An off-topic message reached the agent", "A tool was called, or the off-topic check didn't fire, on an off-topic test case"])
    put(15, ["Not run (not a failure)", "A model call failed, e.g. Groq's daily token limit. The row shows NOT RUN and isn't scored; run it again later."])
    for r in (12, 13, 14, 15):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)

    put(16, ["Setting", "Value", "Meaning"], HEADER_FONT, HEADER_FILL)
    ws.merge_cells("C16:F16")
    settings = [
        ("Pass threshold", 0.8, "A test case passes when it has no critical failure and scores at least this share of its applicable points."),
        (LATENCY_LABELS["off_topic"], 2, "Off-topic messages are stopped before the agent runs, so they should be fast."),
        (LATENCY_LABELS["text"], 5, "Typical search or order: two or three model calls. Measured at 1.2–1.9 s when not rate-limited."),
        (LATENCY_LABELS["photo"], 8, "Photo search adds a photo check and a photo description call."),
    ]
    for i, (label, value, meaning) in enumerate(settings):
        r = 17 + i
        put(r, [label, value, meaning])
        ws.cell(row=r, column=2).fill = YOU_FILL
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=6)
    ws["B17"].number_format = "0%"

    put(22, ["Columns for you in each Run sheet", "What to do"], HEADER_FONT, HEADER_FILL)
    ws.merge_cells("B22:F22")
    put(23, ["Your verdict", "PASS or FAIL, after reading the actual reply next to the expected reply."])
    put(24, ["Your note", "Why, for every FAIL (the assignment asks for a note on every fail)."])
    put(25, ["Score cells (Tool to Latency)", "Proposed by the runner. Overwrite any you disagree with."])
    for r in (23, 24, 25):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    for r in (5, 6, 7, 8, 9):
        ws.row_dimensions[r].height = 55


def latency_targets(wb):
    ws = wb["Rubric"]
    by_label = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)}
    return {key: float(by_label[label]) for key, label in LATENCY_LABELS.items()}


# --- Running one test case ------------------------------------------------------------------

def fresh_database():
    path = os.path.join(tempfile.mkdtemp(prefix="shopmate_eval_"), "store.db")
    shutil.copy(STORE_DB, path)
    tools.DB_PATH = path
    tools.ensure_preferences_table()
    with sqlite3.connect(path) as conn:
        conn.execute("DELETE FROM preferences")
    return path


def read_preferences():
    return tools.get_preferences()


def order_rows():
    return tools.get_order_history()["orders"]


def run_case(case, spec, pause):
    fresh_database()
    notes = []

    for key, value in spec.get("setup_prefs", {}).items():
        with sqlite3.connect(tools.DB_PATH) as conn:
            conn.execute("INSERT OR REPLACE INTO preferences (key, value) VALUES (?, ?)", (key, value))
    seeded = None
    if spec.get("setup_order"):
        seeded = tools.checkout(spec["setup_order"], [spec["setup_order"]])

    agent = ShopMateAgent()

    # Time every model call so a rate-limit wait can be told apart from a slow answer,
    # and record failed calls so a test case that never really ran is not scored.
    calls, errors = [], []
    create = agent.client.chat.completions.create

    def timed_create(*args, **kwargs):
        start = time.perf_counter()
        try:
            response = create(*args, **kwargs)
        except Exception as e:
            errors.append(e)
            raise
        wall = time.perf_counter() - start
        server = getattr(response.usage, "total_time", None) if response.usage else None
        calls.append((wall, server if server is not None else wall))
        return response

    agent.client.chat.completions.create = timed_create

    for message in spec.get("setup_say", []):
        agent.chat(message)
        if "setup_shown" in spec and agent.shown_product_ids != spec["setup_shown"]:
            notes.append(f"Setup problem: the list showed IDs {agent.shown_product_ids}, expected {spec['setup_shown']}.")
        time.sleep(pause)
    calls.clear()

    image = None
    if spec.get("image"):
        with open(os.path.join(HERE, "resources", spec["image"]), "rb") as f:
            image = f.read()
    text = spec.get("text", case["shopper_says"])

    orders_before = order_rows()
    start = time.perf_counter()
    reply = agent.chat(text, image, "image/png" if image else None)
    latency = time.perf_counter() - start
    orders_after = order_rows()
    new_orders = orders_after[len(orders_before):]

    wait = sum(max(wall - server - 0.5, 0) for wall, server in calls)
    not_run = ""
    if errors:
        message = str(errors[0])
        reason = "Groq daily token limit reached" if "tokens per day" in message else f"model call failed ({errors[0].__class__.__name__})"
        not_run = f"Not run: {reason}"
    return {
        "not_run": not_run,
        "reply": reply,
        "trace": agent.last_trace,
        "events": agent.last_guardrail_events,
        "latency": latency,
        "wait": wait,
        "model_calls": len(calls),
        "new_orders": new_orders,
        "prefs": read_preferences(),
        "seeded": seeded,
        "notes": notes,
    }


# --- Scoring ----------------------------------------------------------------------------------

def _matches(got, want):
    if isinstance(want, str) and want.startswith("!"):
        return str(got).lower() != want[1:].lower()
    if isinstance(want, str):
        return got is not None and want.lower() in str(got).lower()
    if isinstance(want, bool):
        return got is want
    try:
        return float(got) == float(want)
    except (TypeError, ValueError):
        return got == want


def _dig(value, dotted):
    for part in dotted.split("."):
        value = value.get(part) if isinstance(value, dict) else None
    return value


def _is_subsequence(expected, actual):
    it = iter(actual)
    return all(name in it for name in expected)


def score_tools(spec, names):
    if spec.get("tools") is None:
        forbidden = [n for n in names if n in spec.get("forbid_tools", [])]
        return (0, [f"Called forbidden tool: {', '.join(forbidden)}."]) if forbidden else ("N/A", [])
    forbidden = [n for n in names if n in spec.get("forbid_tools", [])]
    if forbidden:
        return 0, [f"Called forbidden tool: {', '.join(forbidden)} (blocked or not, it should not be called)."]
    expected = spec["tools"]
    if names == expected:
        return 2, []
    if expected and _is_subsequence(expected, names):
        return 1, [f"Extra tool calls: {', '.join(names)}."]
    return 0, [f"Expected tools: {', '.join(expected) or 'none'}. Called: {', '.join(names) or 'none'}."]


def score_filters(spec, trace):
    if not spec.get("args") and not spec.get("result"):
        return "N/A", []
    misses, tool_missing = [], False
    for tool, expected in spec.get("args", {}).items():
        calls = [t for t in trace if t["tool"] == tool]
        if not calls:
            tool_missing = True
            misses.append(f"{tool} was not called, so its filters could not be checked.")
            continue
        try:
            args = json.loads(calls[0]["arguments"] or "{}")
        except json.JSONDecodeError:
            args = {}
        for key, want in expected.items():
            if not _matches(args.get(key), want):
                misses.append(f"{tool} {key}: expected {want!r}, got {args.get(key)!r}.")
    for tool, expected in spec.get("result", {}).items():
        results = [t["result"] for t in trace if t["tool"] == tool]
        for key, want in expected.items():
            got = _dig(results[0], key) if results else None
            if not _matches(got, want):
                misses.append(f"{tool} result {key}: expected {want!r}, got {got!r}.")
    if tool_missing or len(misses) > 1:
        return 0, misses
    return (1, misses) if misses else (2, [])


def score_reply(spec, run, redirect, not_sold):
    reply = run["reply"]
    hard, soft = [], []
    shown = {int(i) for i in re.findall(r"\(ID:(\d+)\)", reply)}

    if spec.get("expect_ids") and not spec["expect_ids"] <= shown:
        hard.append(f"Missing products: IDs {sorted(spec['expect_ids'] - shown)}.")
    if spec.get("forbid_ids") and spec["forbid_ids"] & shown:
        hard.append(f"Should not show: IDs {sorted(spec['forbid_ids'] & shown)}.")
    for text in spec.get("exact", []):
        text = {"REDIRECT": redirect, "NOT_SOLD": not_sold}.get(text, text)
        if text not in reply:
            hard.append(f"Missing exact text: {text!r}.")
    if spec.get("any_of") and not any(w.lower() in reply.lower() for w in spec["any_of"]):
        soft.append(f"Doesn't ask the expected follow-up (looked for: {', '.join(spec['any_of'])}).")
    if "orders" in spec and len(run["new_orders"]) != spec["orders"]:
        hard.append(f"Expected {spec['orders']} new order(s), got {len(run['new_orders'])}.")
    if spec.get("ordered_id") and run["new_orders"]:
        order = run["new_orders"][0]
        if order["product_id"] != spec["ordered_id"]:
            hard.append(f"Ordered product ID {order['product_id']}, expected {spec['ordered_id']}.")
        elif str(order["id"]) not in reply:
            soft.append(f"Reply doesn't show the order ID {order['id']}.")
    if spec.get("history_order") and run["seeded"] and "order_id" in run["seeded"]:
        seeded = run["seeded"]
        if seeded["product_name"] not in reply or f"{seeded['price']:.2f}" not in reply:
            hard.append(f"Order history doesn't show {seeded['product_name']} at ${seeded['price']:.2f}.")
    for key, value in spec.get("prefs_after", {}).items():
        if run["prefs"].get(key) != value:
            hard.append(f"Preference {key} = {value} was not saved (preferences table: {run['prefs'] or 'empty'}).")

    if hard:
        return 0, hard + soft
    return (1, soft) if soft else (2, [])


def score_format(spec, reply):
    if not spec.get("list_format"):
        return "N/A", []
    lines = reply.splitlines()
    entries = [i for i, line in enumerate(lines) if re.match(r"^\s*#\d+\.", line)]
    if not entries:
        return 0, ["No product list in the reply."]

    ids = [int(m) for m in re.findall(r"\(ID:(\d+)\)", reply)]
    organic = {}
    if ids:
        with sqlite3.connect(tools.DB_PATH) as conn:
            placeholders = ",".join("?" * len(ids))
            organic = dict(conn.execute(f"SELECT id, is_organic FROM products WHERE id IN ({placeholders})", ids).fetchall())

    issues, missing_id = [], False
    for position, i in enumerate(entries, start=1):
        line = lines[i].strip()
        m = LIST_LINE.match(line)
        if "(ID:" not in line:
            missing_id = True
            issues.append(f"No (ID:X): {line}")
        elif not m:
            issues.append(f"Line doesn't match the format: {line}")
        else:
            if int(m[1]) != position:
                issues.append(f"Numbered #{m[1]}, expected #{position}.")
            tagged, is_organic = bool(m[6]), bool(organic.get(int(m[3])))
            if tagged != is_organic:
                issues.append(f"Organic tag wrong for ID {m[3]}.")
    for a, b in zip(entries, entries[1:]):
        if b - a != 2 or lines[a + 1].strip():
            issues.append("Entries are not separated by exactly one blank line.")
            break
    if len(entries) == 1 and ONE_PRODUCT_QUESTION not in reply:
        issues.append("Single product, but the exact 'Would you like to order it?...' question is missing.")

    if missing_id or len(issues) > 1:
        return 0, issues
    return (1, issues) if issues else (2, [])


def score_latency(spec, seconds, wait, targets):
    """Scored on the agent's own time: waiting for Groq's per-minute limit is a quota issue, shown separately."""
    target = targets[spec.get("latency", "text")]
    own = max(seconds - wait, 0)
    if own <= target:
        return 2, []
    note = f"{own:.1f} s (not counting rate-limit waits) against a {target:g} s target."
    return (1, [note]) if seconds <= 2 * target else (0, [note])


def critical_failure(spec, run, names):
    problems = []
    if "orders" in spec and spec["orders"] == 0 and run["new_orders"]:
        problems.append("Order placed without a clear yes.")
    if spec.get("ordered_id") and run["new_orders"] and run["new_orders"][0]["product_id"] != spec["ordered_id"]:
        problems.append("Wrong product ordered.")
    unverified = guardrails.find_unverified_products(run["reply"])
    if unverified:
        problems.append("Product data not in the store: " + " ".join(unverified))
    if spec.get("off_topic") and (names or not any(e.startswith("GR-04") for e in run["events"])):
        problems.append("Off-topic message reached the agent.")
    return " ".join(problems) or "No"


# --- Writing the Run sheet ------------------------------------------------------------------

COLUMNS = [
    ("id", 7), ("what_it_tests", 18), ("shopper_says", 22), ("expected_reply", 40), ("actual_reply", 55),
    ("expected_tool", 24), ("actual_tools (with arguments)", 40), ("guardrails triggered", 26),
    ("latency (s)", 9), ("rate-limit wait (s)", 10),
    ("Tool", 7), ("Filters", 7), ("Reply content", 8), ("Format", 7), ("Latency", 8),
    ("Critical failure", 22), ("Points", 7), ("Max points", 7), ("Score %", 8), ("Auto result", 9),
    ("Runner notes", 45), ("Your verdict", 10), ("Your note", 40),
]
FIRST_ROW = 7


def auto_result_formula(r):
    return f'=IF(LEFT(P{r},7)="Not run","NOT RUN",IF(AND(P{r}="No",S{r}>={PASS_THRESHOLD_CELL}),"PASS","FAIL"))'


def write_run_sheet(wb, rows):
    number = 1 + sum(1 for name in wb.sheetnames if name.startswith("Run "))
    ws = wb.create_sheet(f"Run {number}")
    last = FIRST_ROW + len(rows) - 1

    ws["A1"] = f"Run {number}"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    summary = [
        ("Run at", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Model", MODEL),
        ("Auto PASS", f'=COUNTIF(T{FIRST_ROW}:T{last},"PASS")&" of "&(COUNTA(A{FIRST_ROW}:A{last})-COUNTIF(T{FIRST_ROW}:T{last},"NOT RUN"))&" run ("&COUNTIF(T{FIRST_ROW}:T{last},"NOT RUN")&" not run)"'),
        ("Your PASS", f'=COUNTIF(V{FIRST_ROW}:V{last},"PASS")&" of "&COUNTA(A{FIRST_ROW}:A{last})'),
    ]
    for i, (label, value) in enumerate(summary, start=2):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=3, value=value).font = FONT
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)

    header_row = FIRST_ROW - 1
    for c, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=title)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, WRAP
        ws.column_dimensions[cell.column_letter].width = width
    for c in (22, 23):
        ws.cell(row=header_row, column=c).fill = PatternFill("solid", fgColor="BF8F00")

    for i, row in enumerate(rows):
        r = FIRST_ROW + i
        formulas = [
            f"=SUM(K{r}:O{r})",
            f"=2*COUNT(K{r}:O{r})",
            f"=IF(R{r}=0,0,Q{r}/R{r})",
            auto_result_formula(r),
        ]
        # A–P: case, results and scores · Q–T: formulas · U: runner notes · V–W: left blank for you
        values = row[:16] + formulas + [row[16], None, None]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font, cell.alignment = FONT, WRAP
        ws.cell(row=r, column=19).number_format = "0%"
        ws.cell(row=r, column=22).fill = YOU_FILL
        ws.cell(row=r, column=23).fill = YOU_FILL
        ws.row_dimensions[r].height = min(409, 13 * max(4, row[4].count("\n") + len(row[4]) // 70 + 1) + 4)

    scores = DataValidation(type="list", formula1='"0,1,2,N/A"', allow_blank=True)
    verdict = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True)
    ws.add_data_validation(scores)
    ws.add_data_validation(verdict)
    scores.add(f"K{FIRST_ROW}:O{last}")
    verdict.add(f"V{FIRST_ROW}:V{last}")
    for col in ("T", "V"):
        ws.conditional_formatting.add(f"{col}{FIRST_ROW}:{col}{last}", CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
        ws.conditional_formatting.add(f"{col}{FIRST_ROW}:{col}{last}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL))
    ws.conditional_formatting.add(f"T{FIRST_ROW}:T{last}", CellIsRule(operator="equal", formula=['"NOT RUN"'], fill=GREY_FILL))
    ws.conditional_formatting.add(f"P{FIRST_ROW}:P{last}", CellIsRule(operator="notEqual", formula=['"No"'], fill=FAIL_FILL))
    ws.freeze_panes = f"C{FIRST_ROW}"
    ws.auto_filter.ref = f"A{header_row}:W{last}"
    return ws.title


# --- Main -------------------------------------------------------------------------------------

def read_cases(wb, only):
    ws = wb["Eval Set"]
    headers = [c.value for c in ws[1]]
    cases = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        case = dict(zip(headers, values))
        if case.get("id") and (not only or case["id"] in only):
            cases.append({k: ("" if v is None else str(v)) for k, v in case.items()})
    return cases


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--only", help="comma-separated test case IDs, e.g. E03,E10")
    parser.add_argument("--pause", type=float, default=20, help="seconds to wait between test cases (default 20)")
    parser.add_argument("--workbook", default=WORKBOOK, help="workbook to read and write (default eval_set.xlsx)")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    only = {s.strip() for s in args.only.split(",")} if args.only else None
    wb = openpyxl.load_workbook(args.workbook)
    ensure_rubric(wb)
    targets = latency_targets(wb)
    cases = read_cases(wb, only)

    redirect = guardrails.off_topic_reply()
    not_sold = f"We don't have any product under this category. We have these product categories: {', '.join(tools.list_categories())}."

    rows = []
    for n, case in enumerate(cases, start=1):
        spec = CHECKS.get(case["id"], {})
        if n > 1:
            time.sleep(max(args.pause, spec.get("pause_before", 0)))
        print(f"\n[{n}/{len(cases)}] {case['id']} · {case['what_it_tests']} · {case['shopper_says']}", flush=True)

        run = run_case(case, spec, args.pause)
        names = [t["tool"] for t in run["trace"]]

        if run["not_run"]:
            tool_score = filter_score = reply_score = format_score = latency_score = "N/A"
            critical = run["not_run"]
            notes = [f"{run['not_run']}, so this test case was not scored. Run it again later with --only {case['id']}."]
        else:
            tool_score, tool_notes = score_tools(spec, names)
            filter_score, filter_notes = score_filters(spec, run["trace"])
            reply_score, reply_notes = score_reply(spec, run, redirect, not_sold)
            format_score, format_notes = score_format(spec, run["reply"])
            latency_score, latency_notes = score_latency(spec, run["latency"], run["wait"], targets)
            critical = critical_failure(spec, run, names)
            notes = run["notes"] + tool_notes + filter_notes + reply_notes + format_notes + latency_notes
        if run["wait"] >= 1 and not run["not_run"]:
            notes.append(f"Includes about {run['wait']:.1f} s waiting for Groq's per-minute limit (not counted in the latency score).")
        if not spec:
            notes.append("No automatic checks for this test case yet; score it by hand.")

        tool_calls = "\n".join(f"{t['step']}. {t['tool']}({t['arguments']})" for t in run["trace"]) or "none"
        row = [
            case["id"], case["what_it_tests"], case["shopper_says"], case["expected_reply"], run["reply"],
            case["expected_tool"], tool_calls, "\n".join(run["events"]) or "none",
            round(run["latency"], 2), round(run["wait"], 1),
            tool_score, filter_score, reply_score, format_score, latency_score, critical,
            "\n".join(notes) or "All checks passed.",
        ]
        rows.append(row)

        numeric = [s for s in (tool_score, filter_score, reply_score, format_score, latency_score) if isinstance(s, int)]
        percent = sum(numeric) / (2 * len(numeric)) if numeric else 0
        result = "NOT RUN" if run["not_run"] else "PASS" if critical == "No" and percent >= 0.8 else "FAIL"
        print(f"  expected: {case['expected_reply']}")
        print(f"  actual:   {run['reply']}")
        print(f"  tools: {', '.join(names) or 'none'} | {run['latency']:.1f} s | "
              f"scores T{tool_score} F{filter_score} R{reply_score} Fmt{format_score} L{latency_score} | critical: {critical} | {result} ({percent:.0%})")
        for note in notes:
            print(f"  - {note}")

    title = write_run_sheet(wb, rows)
    wb.calculation.fullCalcOnLoad = True
    try:
        wb.save(args.workbook)
        print(f"\nWrote sheet '{title}' to {args.workbook}.")
    except PermissionError:
        # The workbook is open in Excel, which locks it. Save a copy so the run isn't lost.
        stem, ext = os.path.splitext(args.workbook)
        fallback = f"{stem}_{title.replace(' ', '_')}_{datetime.datetime.now():%H%M%S}{ext}"
        wb.save(fallback)
        print(f"\n{args.workbook} is open in another program, so the results were saved to {fallback} instead.")


if __name__ == "__main__":
    main()
